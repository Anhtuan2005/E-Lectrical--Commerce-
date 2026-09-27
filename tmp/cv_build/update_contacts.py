"""Update contact hyperlinks in the supplied CVs without rebuilding their layout."""

from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn


EMAIL = "Nguyenztuan58@gmail.com"
LINKEDIN_TEXT = "www.linkedin.com/in/anh-tuan-nguyen-truong"
LINKEDIN_URL = "https://www.linkedin.com/in/anh-tuan-nguyen-truong"


def replace_link(paragraph, old_text, new_text, old_url, new_url):
    matches = []
    for link in paragraph._p.findall(qn("w:hyperlink")):
        text_nodes = list(link.iter(qn("w:t")))
        if "".join(node.text or "" for node in text_nodes) == old_text:
            matches.append((link, text_nodes))
    if len(matches) != 1:
        raise ValueError(f"Expected one hyperlink for {old_text!r}; found {len(matches)}")

    link, text_nodes = matches[0]
    relationship = paragraph.part.rels[link.get(qn("r:id"))]
    if relationship.reltype != RT.HYPERLINK or relationship.target_ref != old_url:
        raise ValueError(f"Unexpected hyperlink target: {relationship.target_ref}")
    text_nodes[0].text = new_text
    for node in text_nodes[1:]:
        node.text = ""
    relationship._target = new_url


for name in ("CV_BackEnd_NguyenTruongAnhTuan", "CV_AI_NguyenTruongAnhTuan"):
    source = Path("E:/") / f"{name}.docx"
    output = Path("E:/") / f"{name}_Updated.docx"
    document = Document(source)
    original_paragraphs = [paragraph.text for paragraph in document.paragraphs]

    replace_link(
        document.paragraphs[2],
        "nguyenzatuan58@gmail.com",
        EMAIL,
        "mailto:nguyenzatuan58@gmail.com",
        f"mailto:{EMAIL}",
    )
    replace_link(
        document.paragraphs[3],
        "linkedin.com/in/anh-tuan-nguyen-truong-980404338",
        LINKEDIN_TEXT,
        "https://linkedin.com/in/anh-tuan-nguyen-truong-980404338",
        LINKEDIN_URL,
    )
    document.save(output)

    check = Document(output)
    assert EMAIL in check.paragraphs[2].text
    assert LINKEDIN_TEXT in check.paragraphs[3].text
    for index, (before, after) in enumerate(zip(original_paragraphs, (p.text for p in check.paragraphs))):
        if index not in (2, 3):
            assert before == after, f"Unexpected content change in paragraph {index}"
    targets = [rel.target_ref for rel in check.part.rels.values() if rel.reltype == RT.HYPERLINK]
    assert f"mailto:{EMAIL}" in targets
    assert LINKEDIN_URL in targets
    print(output)
