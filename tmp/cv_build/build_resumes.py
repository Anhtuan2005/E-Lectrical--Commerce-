from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"E:\E-Lectrical--Commerce-")
DOCX_OUT = ROOT / "output" / "docx"
DOCX_OUT.mkdir(parents=True, exist_ok=True)

BLACK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x4A, 0x4A, 0x4A)
LINK = BLACK
FONT = "Arial"


CONTACT = {
    "name": "NGUYEN TRUONG ANH TUAN",
    "phone": "+84 348 021 904",
    "email": "nguyenzatuan58@gmail.com",
    "linkedin_text": "linkedin.com/in/anh-tuan-nguyen-truong-980404338",
    "linkedin_url": "https://linkedin.com/in/anh-tuan-nguyen-truong-980404338",
    "github_text": "github.com/Anhtuan2005",
    "github_url": "https://github.com/Anhtuan2005",
}


BACKEND = {
    "filename": "Nguyen_Truong_Anh_Tuan_Backend_CV.docx",
    "target": "SOFTWARE ENGINEERING STUDENT | .NET | BACKEND | PYTHON",
    "profile": (
        "Software Engineering student building production-style backend applications with ASP.NET Core and Python. "
        "Hands-on experience with REST APIs, SQL data modeling, secure authentication, payment and order workflows, "
        "background processing, structured logging, and automated testing. Seeking backend or .NET software engineering opportunities."
    ),
    "skills": [
        ("Languages", "C#, Python, JavaScript, SQL, HTML5, CSS3"),
        ("Frameworks & Data", "ASP.NET Core 8 MVC, Entity Framework Core, FastAPI, SQL Server, SQLite, Razor, REST APIs, AJAX"),
        ("AI & Tools", "YOLO11, ByteTrack, InsightFace, OpenCV, pytest, Git, Serilog, VNPay"),
        ("Engineering", "API design, authentication, automated testing, logging, data modeling, background processing"),
    ],
    "projects": [
        {
            "title": "TECHVORA | FULL-STACK E-COMMERCE APPLICATION",
            "dates": "02/2026 - 04/2026",
            "stack": "ASP.NET Core 8 MVC | EF Core | SQL Server | JavaScript",
            "repo_text": "github.com/Anhtuan2005/E-Lectrical--Commerce-",
            "repo_url": "https://github.com/Anhtuan2005/E-Lectrical--Commerce-",
            "bullets": [
                "Independently designed and built a Vietnamese e-commerce application using ASP.NET Core 8 MVC, Razor, EF Core, SQL Server and vanilla JavaScript/CSS.",
                "Delivered 10+ customer workflows spanning search, filtering, pagination, AJAX cart and wishlist, accounts, vouchers, checkout and order history.",
                "Integrated 2 payment methods - COD and VNPay sandbox - with webhook validation and consistent payment-to-order lifecycle handling.",
                "Centralized 8 administration domains covering revenue analytics, catalog, inventory audit logs, users, promotions, reviews, shipping and order management.",
                "Applied 3 protection layers across account, checkout and administrative flows: ASP.NET Core Identity, rate limiting and server-side validation.",
                "Instrumented 3 critical event categories - payment callbacks, operations and application errors - with structured Serilog logging.",
            ],
        },
        {
            "title": "SCT CAMERA | REAL-TIME AI SURVEILLANCE SYSTEM",
            "dates": "04/2026 - 08/2026",
            "stack": "Python | FastAPI | YOLO11 | ByteTrack | InsightFace | OpenCV",
            "repo_text": "github.com/Anhtuan2005/SCT_Camera",
            "repo_url": "https://github.com/Anhtuan2005/SCT_Camera",
            "bullets": [
                "Independently designed and built a multi-camera pipeline for 3 source types: webcam, local video and RTSP.",
                "Architected isolated per-camera workers around YOLO11 detection and ByteTrack tracking, exposed through a FastAPI MJPEG dashboard.",
                "Engineered 2 identity-stability mechanisms - per-track face caching and camera motion compensation - for more stable tracking.",
                "Separated InsightFace CPU inference from the YOLO GPU workload to reduce resource contention during multi-camera processing.",
                "Implemented 6 alert workflows: intrusion, loitering, line crossing, falls, unknown persons and suspicious theft or missing assets.",
                "Delivered 20+ authenticated dashboard and API routes for camera, ROI, counting-line, alert-history and YAML configuration workflows.",
                "Hardened 2 notification channels - Telegram and Discord - with 3 delivery safeguards: cooldown, deduplication and retry; validated core workflows with 180 passing automated tests.",
                "Automated 3 operational safeguards - SQLite migrations, persistent logging and graceful shutdown - to release camera resources and background alert workers safely.",
            ],
        },
    ],
}


COMPUTER_VISION = {
    "filename": "Nguyen_Truong_Anh_Tuan_Computer_Vision_CV.docx",
    "target": "SOFTWARE ENGINEERING STUDENT | PYTHON | COMPUTER VISION",
    "profile": (
        "Software Engineering student focused on real-time computer vision systems with Python, YOLO11, ByteTrack, "
        "InsightFace, OpenCV and FastAPI. Hands-on experience with multi-camera pipelines, identity stabilization, "
        "CPU/GPU workload separation, alert workflows and automated testing. Seeking computer vision or applied AI opportunities."
    ),
    "skills": [
        ("Languages", "Python, C#, JavaScript, SQL, HTML5, CSS3"),
        ("Computer Vision", "YOLO11, ByteTrack, InsightFace, OpenCV, object detection, multi-object tracking, face recognition"),
        ("Frameworks & Data", "FastAPI, ASP.NET Core 8 MVC, Entity Framework Core, SQL Server, SQLite, REST APIs"),
        ("Engineering", "Real-time video processing, multi-camera pipelines, automated testing, logging, API design"),
    ],
    "projects": [
        {
            "title": "SCT CAMERA | REAL-TIME AI SURVEILLANCE SYSTEM",
            "dates": "04/2026 - 08/2026",
            "stack": "Python | FastAPI | YOLO11 | ByteTrack | InsightFace | OpenCV",
            "repo_text": "github.com/Anhtuan2005/SCT_Camera",
            "repo_url": "https://github.com/Anhtuan2005/SCT_Camera",
            "bullets": [
                "Independently designed and built a multi-camera pipeline for 3 source types: webcam, local video and RTSP.",
                "Architected isolated per-camera workers around YOLO11 detection and ByteTrack tracking, exposed through a FastAPI MJPEG dashboard.",
                "Engineered 2 identity-stability mechanisms - per-track face caching and camera motion compensation - for more stable tracking.",
                "Separated InsightFace CPU inference from the YOLO GPU workload to reduce resource contention during multi-camera processing.",
                "Implemented 6 alert workflows: intrusion, loitering, line crossing, falls, unknown persons and suspicious theft or missing assets.",
                "Delivered 20+ authenticated dashboard and API routes for camera, ROI, counting-line, alert-history and YAML configuration workflows.",
                "Hardened 2 notification channels - Telegram and Discord - with 3 delivery safeguards: cooldown, deduplication and retry; validated core workflows with 180 passing automated tests.",
                "Automated 3 operational safeguards - SQLite migrations, persistent logging and graceful shutdown - to release camera resources and background alert workers safely.",
            ],
        },
        {
            "title": "TECHVORA | FULL-STACK E-COMMERCE APPLICATION",
            "dates": "02/2026 - 04/2026",
            "stack": "ASP.NET Core 8 MVC | EF Core | SQL Server | JavaScript",
            "repo_text": "github.com/Anhtuan2005/E-Lectrical--Commerce-",
            "repo_url": "https://github.com/Anhtuan2005/E-Lectrical--Commerce-",
            "bullets": [
                "Independently designed and built a Vietnamese e-commerce application using ASP.NET Core 8 MVC, Razor, EF Core, SQL Server and vanilla JavaScript/CSS.",
                "Delivered 10+ customer workflows spanning search, filtering, pagination, AJAX cart and wishlist, accounts, vouchers, checkout and order history.",
                "Integrated 2 payment methods - COD and VNPay sandbox - with webhook validation and consistent payment-to-order lifecycle handling.",
                "Centralized 8 administration domains covering revenue analytics, catalog, inventory audit logs, users, promotions, reviews, shipping and order management.",
                "Applied 3 protection layers across account, checkout and administrative flows: ASP.NET Core Identity, rate limiting and server-side validation.",
                "Instrumented 3 critical event categories - payment callbacks, operations and application errors - with structured Serilog logging.",
            ],
        },
    ],
}


def set_run_font(run, size=None, bold=None, italic=None, color=BLACK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.color.rgb = color
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def set_cellless_spacing(paragraph, before=0, after=0, line=1.0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    return paragraph


def add_hyperlink(paragraph, text, url, size=7.0, color=LINK, underline=False):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), FONT)
    r_fonts.set(qn("w:hAnsi"), FONT)
    r_pr.append(r_fonts)
    r_color = OxmlElement("w:color")
    r_color.set(qn("w:val"), str(color))
    r_pr.append(r_color)
    r_size = OxmlElement("w:sz")
    r_size.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(r_size)
    r_size_cs = OxmlElement("w:szCs")
    r_size_cs.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(r_size_cs)
    underline_node = OxmlElement("w:u")
    underline_node.set(qn("w:val"), "single" if underline else "none")
    r_pr.append(underline_node)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_right_tab(paragraph, position_inches=7.34):
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(position_inches), WD_TAB_ALIGNMENT.RIGHT
    )


def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(9.0)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    title = styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    title.font.size = Pt(21.5)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(0)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    heading = styles["Heading 1"]
    heading.font.name = FONT
    heading._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    heading.font.size = Pt(11.0)
    heading.font.bold = True
    heading.font.color.rgb = BLACK
    heading.paragraph_format.space_before = Pt(6)
    heading.paragraph_format.space_after = Pt(1.8)
    heading.paragraph_format.keep_with_next = True

    if "Resume Subtitle" not in styles:
        subtitle = styles.add_style("Resume Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        subtitle = styles["Resume Subtitle"]
    subtitle.font.name = FONT
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    subtitle.font.size = Pt(10.0)
    subtitle.font.bold = True
    subtitle.font.color.rgb = BLACK
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(0.3)


def add_header(doc, target):
    p = doc.add_paragraph(style="Title")
    set_run_font(p.add_run(CONTACT["name"]), 21.5, bold=True)

    p = doc.add_paragraph(style="Resume Subtitle")
    set_run_font(p.add_run(target), 10.0, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cellless_spacing(p, after=0.8, line=1.0)
    set_run_font(p.add_run(f"Ho Chi Minh City, Vietnam | {CONTACT['phone']} | "), 9.0, color=BLACK)
    add_hyperlink(p, CONTACT["email"], f"mailto:{CONTACT['email']}", size=9.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_cellless_spacing(p, after=1.4, line=1.0)
    add_hyperlink(p, CONTACT["linkedin_text"], CONTACT["linkedin_url"], size=8.5)
    set_run_font(p.add_run(" | "), 8.5, color=BLACK)
    add_hyperlink(p, CONTACT["github_text"], CONTACT["github_url"], size=8.5)


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    spacing = {
        "PROFESSIONAL PROFILE": (7.5, 3.4),
        "TECHNICAL SKILLS": (7.5, 2.8),
        "SELECTED PROJECTS": (6.7, 3.4),
        "EDUCATION & CREDENTIALS": (7.6, 3.1),
    }
    before, after = spacing[text]
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    set_run_font(p.add_run(text), 11.0, bold=True)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "A6A6A6")
    borders.append(bottom)
    p_pr.append(borders)
    return p


def add_profile(doc, text):
    add_section_heading(doc, "PROFESSIONAL PROFILE")
    p = doc.add_paragraph()
    set_cellless_spacing(p, after=0.8, line=1.0)
    set_run_font(p.add_run(text), 9.0)


def add_skills(doc, rows):
    add_section_heading(doc, "TECHNICAL SKILLS")
    for label, value in rows:
        p = doc.add_paragraph()
        set_cellless_spacing(p, after=0, line=1.0)
        set_run_font(p.add_run(f"{label}: "), 9.0, bold=True)
        set_run_font(p.add_run(value), 9.0, color=MUTED)


def add_bullet(doc, text):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.left_indent = Inches(0.14)
    fmt.first_line_indent = Inches(-0.10)
    fmt.right_indent = Inches(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0.5)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.keep_together = True
    set_run_font(p.add_run("- "), 9.0)
    set_run_font(p.add_run(text), 9.0)
    return p


def add_projects(doc, projects):
    add_section_heading(doc, "SELECTED PROJECTS")
    for index, project in enumerate(projects):
        p = doc.add_paragraph()
        set_cellless_spacing(p, before=6.5 if index else 0, after=1.45, line=1.0)
        p.paragraph_format.keep_with_next = True
        set_right_tab(p)
        set_run_font(p.add_run(project["title"]), 9.5, bold=True)
        set_run_font(p.add_run("\t" + project["dates"]), 9.0, bold=True)

        p = doc.add_paragraph()
        set_cellless_spacing(p, after=0.9, line=1.0)
        p.paragraph_format.keep_with_next = True
        set_run_font(p.add_run("STACK "), 8.0, bold=True)
        set_run_font(p.add_run(project["stack"]), 8.0, color=MUTED)

        p = doc.add_paragraph()
        set_cellless_spacing(p, after=1.55, line=1.0)
        p.paragraph_format.keep_with_next = True
        set_run_font(p.add_run("GITHUB "), 8.0, bold=True)
        add_hyperlink(p, project["repo_text"], project["repo_url"], size=8.0)

        for bullet in project["bullets"]:
            add_bullet(doc, bullet)


def add_education(doc):
    add_section_heading(doc, "EDUCATION & CREDENTIALS")

    p = doc.add_paragraph()
    set_cellless_spacing(p, after=0.7, line=1.0)
    set_right_tab(p)
    set_run_font(
        p.add_run("Ho Chi Minh City University of Foreign Languages - Information Technology"),
        9.0,
        bold=True,
    )
    set_run_font(p.add_run("\t06/2023 - 04/2027"), 8.5, bold=True)

    p = doc.add_paragraph()
    set_cellless_spacing(p, after=0.7, line=1.0)
    set_run_font(p.add_run("Bachelor of Information Technology"), 9.0, italic=True)
    set_run_font(p.add_run(" | GPA: "), 9.0)
    set_run_font(p.add_run("3.16/4.0"), 9.0, bold=True)
    set_run_font(p.add_run(" | IELTS: "), 9.0)
    set_run_font(p.add_run("6.5"), 9.0, bold=True)

    p = doc.add_paragraph()
    set_cellless_spacing(p, after=0, line=1.0)
    set_run_font(p.add_run("Relevant Coursework: "), 9.0, bold=True)
    set_run_font(
        p.add_run("Software Engineering, Web Development, Database Management, Computer Networks"),
        9.0,
        color=BLACK,
    )


def build_resume(config):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(0.665)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    section.header_distance = Inches(0.1)
    section.footer_distance = Inches(0.1)

    configure_styles(doc)
    add_header(doc, config["target"])
    add_profile(doc, config["profile"])
    add_skills(doc, config["skills"])
    add_projects(doc, config["projects"])
    add_education(doc)

    doc.core_properties.title = f"{CONTACT['name']} - {config['target']}"
    doc.core_properties.subject = "Resume"
    doc.core_properties.author = CONTACT["name"].title()
    doc.core_properties.keywords = "resume, software engineering, backend, computer vision"

    output_path = DOCX_OUT / config["filename"]
    doc.save(output_path)
    return output_path


if __name__ == "__main__":
    outputs = [build_resume(BACKEND), build_resume(COMPUTER_VISION)]
    for output in outputs:
        print(output)
